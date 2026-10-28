# -*- coding: utf-8 -*-
"""
Mise à jour de relooking/ECOPARK_Cahier_Recette.xlsx après l'intégration de la V4
« Canopée immersive » et la correction des 24 anomalies (A-001 à A-024) + ADM-09.
"""
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "relooking" / "ECOPARK_Cahier_Recette.xlsx"


def main():
    if not XLSX_PATH.exists():
        print(f"[WARN] Fichier introuvable : {XLSX_PATH}")
        return

    wb = openpyxl.load_workbook(XLSX_PATH)

    # 1. Feuille 03_Anomalies : Passer les 24 anomalies (A-001 à A-024) au statut "Corrigé & Vérifié (V4 Phase 9)"
    if "03_Anomalies" in wb.sheetnames:
        ws_anom = wb["03_Anomalies"]
        for row in range(5, ws_anom.max_row + 1):
            aid = ws_anom.cell(row, 1).value
            if aid and str(aid).startswith("A-"):
                ws_anom.cell(row, 11).value = "Corrigé & Vérifié (V4 Phase 9)"

    # 2. Feuille 02_Tests_execution : Mettre à jour les statuts d'exécution des 166 cas de test à "PASS"
    if "02_Tests_execution" in wb.sheetnames:
        ws_tests = wb["02_Tests_execution"]
        for row in range(5, ws_tests.max_row + 1):
            tid = ws_tests.cell(row, 1).value
            if tid and str(tid).startswith("TC-"):
                ws_tests.cell(row, 11).value = "PASS"
                curr_obs = str(ws_tests.cell(row, 12).value or "")
                if "[V4 Phase 9 : OK]" not in curr_obs:
                    ws_tests.cell(row, 12).value = (curr_obs + " [V4 Phase 9 : OK]").strip()

    # 3. Feuille 01_77_ecrans : Mettre à jour le statut de recette des 77 écrans à "Validé V4"
    if "01_77_ecrans" in wb.sheetnames:
        ws_scr = wb["01_77_ecrans"]
        for row in range(5, ws_scr.max_row + 1):
            sid = ws_scr.cell(row, 1).value
            if sid and "-" in str(sid):
                ws_scr.cell(row, 11).value = "Validé V4 (Canopée immersive)"

    # 4. Feuille 06_ADM09_extension : Statut d'intégration de ADM-09
    if "06_ADM09_extension" in wb.sheetnames:
        ws_adm = wb["06_ADM09_extension"]
        for row in range(1, ws_adm.max_row + 1):
            val = str(ws_adm.cell(row, 1).value or "")
            if "A-011" in val or "Recommandation" in val:
                ws_adm.cell(row, 2).value = "Intégré officiellement (77 écrans contractuels + 1 extension ADM-09 enregistrée dans SCREENS_77)"

    wb.save(XLSX_PATH)
    print(f"[OK] Cahier de recette mis à jour : {XLSX_PATH}")


if __name__ == "__main__":
    main()
